from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

from .config import settings

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx"}


def _normalize_metadata(path: Path, **extra) -> dict:
    return {
        "source": path.name,
        "source_path": str(path.resolve()),
        **extra,
    }


def load_pdf(path: Path) -> list[Document]:
    reader = PdfReader(str(path))
    docs: list[Document] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            docs.append(
                Document(
                    page_content=text,
                    metadata=_normalize_metadata(path, page=index),
                )
            )
    return docs


def load_text(path: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return [Document(page_content=text, metadata=_normalize_metadata(path))]


def load_markdown(path: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    heading_re = re.compile(r"^(#{1,6})\s+(.+)$", re.MULTILINE)
    matches = list(heading_re.finditer(text))

    if not matches:
        return [Document(page_content=text, metadata=_normalize_metadata(path))]

    docs: list[Document] = []
    for i, match in enumerate(matches):
        start = match.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        section_title = match.group(2).strip()
        content = text[start:end].strip()
        if content:
            docs.append(
                Document(
                    page_content=f"{section_title}\n\n{content}",
                    metadata=_normalize_metadata(path, section=section_title),
                )
            )
    return docs


def load_docx(path: Path) -> list[Document]:
    docx = DocxDocument(str(path))
    docs: list[Document] = []
    current_heading = "Document"
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer
        text = "\n".join(buffer).strip()
        if text:
            docs.append(
                Document(
                    page_content=text,
                    metadata=_normalize_metadata(path, section=current_heading),
                )
            )
        buffer = []

    for paragraph in docx.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
        if style_name.startswith("heading"):
            flush()
            current_heading = text
        else:
            buffer.append(text)
    flush()
    return docs


def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return load_pdf(path)
    if suffix == ".md":
        return load_markdown(path)
    if suffix == ".txt":
        return load_text(path)
    if suffix == ".docx":
        return load_docx(path)
    raise ValueError(f"Unsupported file type: {path.suffix}")


def discover_documents(directories: Iterable[Path] | None = None) -> list[Path]:
    directories = list(directories or settings.all_document_dirs)
    paths: list[Path] = []
    for directory in directories:
        if not directory.exists():
            continue
        for path in directory.rglob("*"):
            if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
                paths.append(path)
    return sorted(paths)


def load_and_split_documents(paths: list[Path] | None = None) -> list[Document]:
    paths = paths or discover_documents()
    raw_docs: list[Document] = []
    for path in paths:
        raw_docs.extend(load_file(path))

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_documents(raw_docs)

    for idx, chunk in enumerate(chunks):
        chunk.metadata["chunk_id"] = idx
    return chunks
